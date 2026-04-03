from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from sbl_audit_gui_v_2000 import (
    SYSTEM_COLUMNS,
    AuditWorkbookService,
    AuditEngine,
)


def prepare_vm_case_workbook(source_path: Path, output_path: Path) -> None:
    wb = load_workbook(source_path)
    ws = wb.active

    headers = [str(ws.cell(1, c).value).strip() if ws.cell(1, c).value is not None else "" for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}

    target_col = idx.get("ArcGIS_WebAdaptor")
    if not target_col:
        raise ValueError("Missing ArcGIS_WebAdaptor column in test workbook header row")

    version_col = idx.get("VERSION LOCATIONS")
    if not version_col:
        raise ValueError("Missing VERSION LOCATIONS column in test workbook header row")

    first_data_row = None
    for row_idx in range(2, ws.max_row + 1):
        target_cell = ws.cell(row_idx, target_col)
        software_value = ws.cell(row_idx, 1).value
        if isinstance(target_cell, MergedCell):
            continue
        if software_value is None or str(software_value).strip() == "":
            continue
        first_data_row = row_idx
        break

    if first_data_row is None:
        raise ValueError("Could not find a writable data row for VM smoke test")

    ws.cell(first_data_row, target_col).value = "X"
    if not ws.cell(first_data_row, version_col).value:
        ws.cell(first_data_row, version_col).value = "PROGRAMS AND FEATURES"

    wb.save(output_path)


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    source = root / "tests" / "testing_sbl.xlsx"
    prepared = root / "tests" / f"_audit_vm_case_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    print(f"VM_TEST_SOURCE={source}")
    print(f"VM_TEST_PREPARED={prepared}")

    prepare_vm_case_workbook(source, prepared)

    workbook_service = AuditWorkbookService(str(prepared))
    workbook_service.detect_header_row()
    workbook_service.build_column_map()

    vm_profile = {
        "targets": {
            name: {
                "vm_name": ("FAKE_VM_01" if name == "ArcGIS_WebAdaptor" else ""),
                "os_type": "windows",
            }
            for name in SYSTEM_COLUMNS
        }
    }

    logs = []

    def logger(msg: str):
        logs.append(msg)
        print(msg)

    engine = AuditEngine(
        workbook_service=workbook_service,
        logger=logger,
        vm_profile=vm_profile,
        vcenter_creds={"server": "", "username": "", "password": ""},
        guest_creds={"username": "", "password": ""},
    )

    try:
        engine.run()
        print("AUDIT_VM_SMOKE_TEST=PASS")
        return 0
    except Exception as exc:
        print(f"AUDIT_VM_SMOKE_TEST=FAIL: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
