"""Template asset management and updates service."""

import json
import re
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill

from config import (
    MASTER_SOFTWARE_LIST_PATH,
    SYSTEM_COLUMNS,
    TEMPLATES_DIR,
)
from models import AuditRow, WorkbookSchema
from services.file_logger import FileLogger
from services.utils import (
    _empty_master_software_list_payload,
    _new_model_bucket,
    auto_fit_columns,
    default_target_columns,
    derive_import_target_columns,
    derive_path_metadata,
    ensure_project_structure,
    extract_path_from_version_location,
    get_master_json_template_baseline_path,
    get_master_json_template_latest_path,
    get_sbl_template_baseline_path,
    get_sbl_template_latest_path,
    infer_build_type,
    is_known_non_target_header,
    normalize_header,
    normalize_model_key,
    normalize_text,
    normalize_version,
    pick_record_value,
    today_str,
    timestamp_str,
)


def _get_sbl_model_from_workbook(sbl_path: str) -> str:
    """Extract SBL model name (first component in workbook, e.g., 'Geospatial Intelligence Foundation')."""
    try:
        # Import here to avoid circular dependency
        from services.workbook_service import AuditWorkbookService
        service = AuditWorkbookService(sbl_path)
        service.detect_header_row()
        service.build_column_map()
        for row in service.iter_audit_rows():
            if row.software_component:
                return row.software_component
    except Exception:
        pass
    return "UNKNOWN"


def _model_to_filename_slug(model_name: str) -> str:
    """Convert model name to filename slug.
    E.g., 'Geospatial Intelligence Foundation' -> 'GEOINT_FD' (acronym from key words).
    """
    if "Geospatial" in model_name and "Intelligence" in model_name:
        return "GEOINT_FD"  # Special case for this known model

    # Generic fallback: first word first 6 chars + last word first 2 chars
    words = model_name.split()
    if len(words) >= 2:
        return f"{words[0][:6].upper()}_{words[-1][:2].upper()}"
    return model_name[:8].upper()


class TemplateAssetService:
    def __init__(self):
        """Initialize the TemplateAssetService instance."""
        ensure_project_structure()

    @staticmethod
    def _find_first_existing(paths: List[Path]) -> Optional[Path]:
        """Internal helper for find first existing."""
        for candidate in paths:
            if candidate.exists():
                return candidate
        return None

    @staticmethod
    def _blank_master_payload(components: Dict[str, str], sbl_model: Optional[str] = None) -> Dict[str, Any]:
        """Internal helper for blank master payload."""
        model_key = normalize_model_key(sbl_model)
        software_components = {}
        for name, version_location in components.items():
            meta = derive_path_metadata(version_location)
            software_components[name] = {
                "generic_name": name,
                "registry_name": "",
                "path": meta["path"],
                "coded_path": meta["coded_path"],
                "verification_source": meta["verification_source"],
                "path_last_verified_date": "",
                "notes": "",
                "tracked_in_vms": [],
            }

        payload = {
            "SBL_models": {
                "GENERAL": _new_model_bucket(),
                model_key: {
                    "software_components": software_components,
                    "vm_components": {
                        vm_name: {
                            "vm_name": "",
                            "os_type": "windows",
                            "tracked_software": [],
                        }
                        for vm_name in SYSTEM_COLUMNS
                    },
                },
            },
            "updated_at": datetime.now().isoformat(timespec="seconds"),
        }
        return payload

    def _extract_components_from_sbl(self, sbl_path: str) -> Dict[str, str]:
        """Internal helper for extract components from sbl."""
        # Import here to avoid circular dependency
        from services.workbook_service import AuditWorkbookService
        service = AuditWorkbookService(sbl_path)
        service.detect_header_row()
        service.build_column_map()
        components: Dict[str, str] = {}
        for row in service.iter_audit_rows():
            if row.software_component and row.software_component not in components:
                components[row.software_component] = row.version_locations
        return components

    def _write_blank_master_from_sbl(self, sbl_path: str, output_json_path: Path, sbl_model: Optional[str] = None) -> str:
        """Internal helper for write blank master from sbl."""
        components = self._extract_components_from_sbl(sbl_path)
        model_key = normalize_model_key(sbl_model)
        if output_json_path.exists():
            payload = json.loads(output_json_path.read_text(encoding="utf-8"))
            if "SBL_models" not in payload or not isinstance(payload.get("SBL_models"), dict):
                payload = self._blank_master_payload({}, sbl_model="GENERAL")
        else:
            payload = self._blank_master_payload({}, sbl_model="GENERAL")

        model_bucket = payload.setdefault("SBL_models", {}).setdefault(model_key, _new_model_bucket())
        model_bucket["software_components"] = self._blank_master_payload(components, sbl_model=model_key)["SBL_models"][model_key]["software_components"]
        payload["updated_at"] = datetime.now().isoformat(timespec="seconds")
        output_json_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        return str(output_json_path)

    def update_baseline_from_sbl(self, sbl_path: str) -> Dict[str, str]:
        """Update baseline templates from SBL and upsert model in master software list template."""
        sbl_model = _get_sbl_model_from_workbook(sbl_path)
        baseline_sbl = get_sbl_template_baseline_path(sbl_model)
        baseline_json = get_master_json_template_baseline_path(sbl_model)

        shutil.copy2(sbl_path, baseline_sbl)
        json_path = self._write_blank_master_from_sbl(sbl_path, baseline_json, sbl_model=sbl_model)
        return {
            "sbl_template": str(baseline_sbl),
            "json_template": json_path,
        }

    def update_latest_sbl(self, sbl_path: str) -> str:
        """Copy the given SBL workbook to the model-specific latest-snapshot path."""
        sbl_model = _get_sbl_model_from_workbook(sbl_path)
        latest_path = get_sbl_template_latest_path(sbl_model)
        shutil.copy2(sbl_path, latest_path)
        return str(latest_path)

    def snapshot_current_master_to_latest(self, sbl_model: Optional[str] = None) -> str:
        """Keep the canonical master software list in JSON folder only."""
        if not MASTER_SOFTWARE_LIST_PATH.exists():
            payload = _empty_master_software_list_payload()
            MASTER_SOFTWARE_LIST_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        return str(MASTER_SOFTWARE_LIST_PATH)

    def get_baseline_sbl_path(self, sbl_model: Optional[str] = None) -> str:
        """Get baseline SBL path (model-specific if provided)."""
        return str(get_sbl_template_baseline_path(sbl_model))

    def resolve_existing_baseline_sbl_path(self, sbl_model: Optional[str] = None) -> Optional[str]:
        """Return an existing baseline template path, trying model-specific then fallbacks."""
        preferred = get_sbl_template_baseline_path(sbl_model)
        default_path = get_sbl_template_baseline_path()
        candidates = [preferred, default_path]
        fallback = self._find_first_existing(candidates)
        if fallback is not None:
            return str(fallback)

        # Final fallback: any baseline template in templates folder.
        any_baseline = sorted(TEMPLATES_DIR.glob("sbl_template_baseline_*.xlsx"))
        if any_baseline:
            return str(any_baseline[0])
        return None

    @staticmethod
    def _format_mtime(path: Path) -> str:
        """Internal helper for format mtime."""
        if not path.exists():
            return "Missing"
        return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")

    def get_template_status(self, sbl_model: Optional[str] = None) -> Dict[str, Dict[str, str]]:
        """Get template status (model-specific if provided)."""
        if sbl_model is None:
            sbl_model = "Geospatial Intelligence Foundation"  # default

        targets = {
            "baseline_sbl": get_sbl_template_baseline_path(sbl_model),
            "latest_sbl": get_sbl_template_latest_path(sbl_model),
            "baseline_master_list": get_master_json_template_baseline_path(sbl_model),
        }
        status: Dict[str, Dict[str, str]] = {}
        for key, path in targets.items():
            status[key] = {
                "path": str(path),
                "exists": "Yes" if path.exists() else "No",
                "updated": self._format_mtime(path),
            }
        status["latest_master_list"] = {
            "path": str(MASTER_SOFTWARE_LIST_PATH),
            "exists": "Yes" if MASTER_SOFTWARE_LIST_PATH.exists() else "No",
            "updated": self._format_mtime(MASTER_SOFTWARE_LIST_PATH),
        }
        return status

    def import_list_to_sbl_workbook(self, input_path: str, output_path: str) -> str:
        """Import list to sbl workbook."""
        ext = Path(input_path).suffix.lower()
        if ext in (".xlsx", ".xlsm"):
            shutil.copy2(input_path, output_path)
            return output_path

        rows: List[Dict[str, Any]] = []
        if ext == ".csv":
            df = pd.read_csv(input_path)
            raw_records = df.fillna("").to_dict(orient="records")
            rows = [
                {str(key): value for key, value in record.items()}
                for record in raw_records
                if isinstance(record, dict)
            ]
        elif ext == ".json":
            payload = json.loads(Path(input_path).read_text(encoding="utf-8"))
            if isinstance(payload, dict):
                candidate_rows = payload.get("rows", [])
                if isinstance(candidate_rows, list):
                    rows = [item for item in candidate_rows if isinstance(item, dict)]
            elif isinstance(payload, list):
                rows = [item for item in payload if isinstance(item, dict)]
        else:
            raise ValueError("Unsupported import format. Use .xlsx, .xlsm, .csv, or .json")

        baseline_template = get_sbl_template_baseline_path()
        if baseline_template.exists():
            # Preserve baseline sheet structure (merged headers, styles, and banner rows)
            # by seeding imports from the canonical template instead of a blank workbook.
            shutil.copy2(baseline_template, output_path)
            wb = load_workbook(output_path)
            ws = wb.active
            if ws is None:
                raise ValueError("Baseline workbook has no active worksheet")

            def set_cell_value(row_idx: int, col_idx: Optional[int], value: str) -> None:
                if not col_idx:
                    return
                cell = ws.cell(row_idx, col_idx)
                if isinstance(cell, MergedCell):
                    return
                cell.value = value

            from services.workbook_service import AuditWorkbookServiceProxy
            proxy = AuditWorkbookServiceProxy(ws)
            col_map, sbl_header, audit_header = proxy.build_column_map()
            header_row = proxy.header_row_index or 1

            software_col = col_map.get("SOFTWARE COMPONENT")
            current_ci_header = next((name for name in col_map.keys() if name.startswith("CURRENT CI VERSION")), "")
            current_ci_col = col_map.get(current_ci_header)
            version_locations_col = col_map.get("VERSION LOCATIONS")
            sbl_col = col_map.get(sbl_header) if sbl_header else None
            audit_col = col_map.get(audit_header) if audit_header else None

            template_target_columns = [name for name in col_map.keys() if not is_known_non_target_header(name)]

            data_start_row = header_row + 1
            if software_col:
                for row_idx in range(header_row + 1, ws.max_row + 1):
                    cell_value = normalize_text(ws.cell(row_idx, software_col).value)
                    if not cell_value:
                        continue
                    if "geospatial intelligence foundation" in cell_value.lower():
                        data_start_row = row_idx + 3
                        break
                    data_start_row = row_idx
                    break

            # Clear existing data rows while preserving template styles and merges.
            for row_idx in range(data_start_row, ws.max_row + 1):
                for col_idx in col_map.values():
                    cell = ws.cell(row_idx, col_idx)
                    if not isinstance(cell, MergedCell):
                        cell.value = None

            for offset, record in enumerate(rows):
                row_idx = data_start_row + offset
                software_component = pick_record_value(record, ["SOFTWARE COMPONENT", "NAME", "COMPONENT"])
                current_ci = pick_record_value(record, ["CURRENT CI VERSION", "CURRENT VERSION", "EXPECTED VERSION", "VERSION"])
                sbl_build = pick_record_value(record, ["SBL BUILD", "SBL BUILD VERSION", "BUILD VERSION", "BASELINE"])
                audit_value = pick_record_value(record, ["AUDIT", "AUDIT VALUE"])
                version_locations = pick_record_value(record, ["VERSION LOCATIONS", "VERSION LOCATION", "PATH", "LOCATION", "RULE"])
                target_vms = record.get("target_vms", {}) if isinstance(record.get("target_vms", {}), dict) else {}

                set_cell_value(row_idx, software_col, software_component)
                set_cell_value(row_idx, current_ci_col, current_ci)
                set_cell_value(row_idx, sbl_col, sbl_build)
                set_cell_value(row_idx, audit_col, audit_value)
                set_cell_value(row_idx, version_locations_col, version_locations)

                for target_name in template_target_columns:
                    target_col = col_map.get(target_name)
                    if not target_col:
                        continue
                    if target_name in target_vms:
                        target_value = normalize_text(target_vms.get(target_name, ""))
                    else:
                        target_value = pick_record_value(record, [target_name])
                    set_cell_value(row_idx, target_col, target_value)

            wb.save(output_path)
            return output_path

        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet("Imported SBL")
        ws.title = "Imported SBL"
        target_columns = derive_import_target_columns(rows)
        headers = [
            "SOFTWARE COMPONENT",
            "CURRENT CI VERSION",
            "VERSION LOCATIONS",
            *target_columns,
            "SBL BUILD TEMPLATE",
            "AUDIT TEMPLATE",
        ]
        ws.append(headers)

        for record in rows:
            software_component = pick_record_value(record, ["SOFTWARE COMPONENT", "NAME", "COMPONENT"])
            current_ci = pick_record_value(record, ["CURRENT CI VERSION", "CURRENT VERSION", "EXPECTED VERSION", "VERSION"])
            sbl_build = pick_record_value(record, ["SBL BUILD", "SBL BUILD VERSION", "BUILD VERSION", "BASELINE"])
            audit_value = pick_record_value(record, ["AUDIT", "AUDIT VALUE"])
            version_locations = pick_record_value(record, ["VERSION LOCATIONS", "VERSION LOCATION", "PATH", "LOCATION", "RULE"])
            target_vms = record.get("target_vms", {}) if isinstance(record.get("target_vms", {}), dict) else {}
            row_values = [
                software_component,
                current_ci,
                version_locations,
            ]
            for target_name in target_columns:
                if target_name in target_vms:
                    row_values.append(normalize_text(target_vms.get(target_name, "")))
                else:
                    row_values.append(pick_record_value(record, [target_name]))
            row_values.extend([
                sbl_build,
                audit_value,
            ])
            ws.append(row_values)

        auto_fit_columns(ws)
        wb.save(output_path)
        return output_path
