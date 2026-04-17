"""
AuditMatic Configuration Module

Constants, paths, and default settings used throughout the application.
"""

from pathlib import Path
from openpyxl.styles import PatternFill

# Abbreviations reference:
# SBL - Software Build List
# TG - Tool Generated (files)
# UP - User Provided (files)

# =====================================================================
# Application Information
# =====================================================================
APP_TITLE = "Audit Tool v2"
APP_GEOMETRY = "1020x760"
LOCAL_SENTINEL = "__LOCAL__"

# =====================================================================
# Directory Paths
# =====================================================================
APP_DIR = Path(__file__).resolve().parent
PROJECT_DIR = APP_DIR

AUDIT_RESULTS_DIR = PROJECT_DIR / "Audit Results"
AUDIT_CHECKLIST_DIR = PROJECT_DIR / "Audit Checklist"
JSON_DIR = PROJECT_DIR / "JSON"
JSON_RESULTS_DIR = JSON_DIR / "json_result"
JSON_CHECKLIST_DIR = JSON_DIR / "json_checklist"
JSON_REGISTRY_SNAPSHOTS_DIR = JSON_DIR / "registry_snapshots"
JSON_SCAN_JOBS_DIR = JSON_DIR / "scan_jobs"
PROFILES_DIR = PROJECT_DIR / "profiles" / "vm_profiles"
LOGS_DIR = PROJECT_DIR / "logs"
TESTS_DIR = PROJECT_DIR / "tests"
TEMPLATES_DIR = PROJECT_DIR / "templates"

# =====================================================================
# Master File Paths
# =====================================================================
MASTER_SOFTWARE_LIST_PATH = JSON_DIR / "master_software_list.json"

# =====================================================================
# Template Path Defaults (Backward Compatibility)
# =====================================================================
# These are populated by lazy-loading functions in utils.py
SBL_TEMPLATE_BASELINE_PATH = None  # Lazy-loaded
SBL_TEMPLATE_LATEST_PATH = None  # Lazy-loaded
MASTER_JSON_TEMPLATE_BASELINE_PATH = None  # Lazy-loaded
MASTER_JSON_TEMPLATE_LATEST_PATH = None  # Lazy-loaded

# =====================================================================
# Column and Header Definitions
# =====================================================================
REQ_HEADERS = {"SOFTWARE COMPONENT", "CURRENT CI VERSION", "VERSION LOCATIONS"}

# Default VM target columns (system-wide defaults for backward compatibility)
SYSTEM_COLUMNS = [
    "ArcGIS_WebAdaptor",
    "ArcGIS_Portal",
    "ArcGIS_HostingServer",
    "ArcGIS_DS1",
    "ArcGIS_DS2",
    "GCS_Management",
]

# Valid build type options for profiles
BUILD_TYPE_OPTIONS = ["unknown", "baseline", "latest", "custom"]

# =====================================================================
# Excel Cell Styling (Colors for Pass/Fail/Warn)
# =====================================================================
PASS_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
WARN_FILL = PatternFill(fill_type="solid", fgColor="FFEB9C")

# =====================================================================
# Optional Dependency Availability Flags
# =====================================================================
# These will be set by services that require optional imports
PYVMOMI_AVAILABLE = False
PARAMIKO_AVAILABLE = False
CRYPTO_AVAILABLE = False
